"""
Upgrade 17: the weekly pipeline in three commands (first run on GW6, part3 added for GW7).

  python weekly.py part1 "2026-27 after GW6"
      reads the scores from your model's Results sheet, downloads Understat xG, bookmaker odds and
      FPL availability, loads results / xG / odds / availability into prem_sims and prints the checks.
      Safe to rerun: staging is emptied first and every write overwrites the same rows.

  (in Excel: paste availability.csv into Availability, set NextGW, refresh the xG table, F9, Ctrl+S, close)

  python weekly.py part2 "2026-27 after GW6"
      reads the new ratings (Team Strength updated A, J, K) and the 1,000 simulations straight from
      the saved model, loads the run, then runs Dixon-Coles, the benchmarks, the market model,
      the uncertainty sims and the ensemble for that run, and prints the checks.
      Refuses a label that is already loaded (weekly guide Troubleshooting F1).

  python weekly.py part3 "2026-27 after GW6"
      the Upgrade 15 weekly steps for the next gameweek: GW+1 referees (writes referees_next.csv the
      first time and stops so you can type the names; run part3 again once filled in), the 15.5a
      expected-goals insert, player_data update, then player_minutes, player_rates, team_counts,
      player_markets and match_markets 'next', then every check with OK / CHECK.
      Safe to rerun: every step replaces its own rows for this run.
      Add --skip-referees to carry on without referees (factor 1, "referee TBC").

  python weekly.py check "2026-27 after GW6"
      just the checks for a run (part1, part2 and part3 tables).

  python weekly.py odds
      downloads bookmaker odds for the upcoming games (football-data.co.uk fixtures.csv: 1X2 and
      over/under 2.5, Bet365, market average and best price) into next_odds / v_next_odds, which the
      dashboard's H2H value section reads. part3 runs it too; rerun it on Friday for fresher prices.
      It also logs every next-gameweek price with the model's chance and edge into value_log;
      part1 later fills in the closing odds and the result, so v_value_summary shows whether the
      flagged prices beat the closing line over the season.

The database password is read from the PREM_DB_URL environment variable (Upgrade 16).
"""
import csv
import gzip
import json
import os
import subprocess
import sys
import urllib.request
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

import dc_model

# ============================== CONFIG ==============================
MODEL_PATH = r"C:\Users\YOUR_NAME\OneDrive\Prem Project\prem simulation 26-27 EXPORT.xlsx"
EXPORT_ROOT = r"C:\Users\YOUR_NAME\OneDrive\Prem Project\exports"
SEASON = "2026-27"
UNDERSTAT_URL = "https://understat.com/getLeagueData/EPL/2026"
ODDS_URL = "https://www.football-data.co.uk/mmz4281/2627/E0.csv"
NEXT_ODDS_URL = "https://www.football-data.co.uk/fixtures.csv"   # upcoming games' odds (updated Fri / Tue)
# value log (Upgrade 17): same rule as H2H's "Value vs bookmaker" section
VALUE_SOURCE = "b365"        # whose prices are logged and compared with the close: b365, avg or max
VALUE_MODEL_WEIGHT = 0.5     # blended chance = weight x model + (1 - weight) x bookie's fair chance
VALUE_FLAG_EDGE = 0.05       # VALUE when blended chance x odds - 1 is at least this
# Understat full names -> names your team_aliases table knows (same list as make_weekly_csvs.py)
UNDERSTAT_NAMES = {
    "Manchester City": "Man City",
    "Manchester United": "Man United",
    "Newcastle United": "Newcastle",
    "Nottingham Forest": "Nott'm Forest",
    "Wolverhampton Wanderers": "Wolves",
}
# ====================================================================

ENGINE = create_engine(os.environ.get("PREM_DB_URL", dc_model.DB_URL))


# ---------------------------------------------------------------- helpers
def fetch(url, extra=None):
    h = {"User-Agent": "Mozilla/5.0"}
    h.update(extra or {})
    with urllib.request.urlopen(urllib.request.Request(url, headers=h), timeout=60) as r:
        data = r.read()
    if data[:2] == b"\x1f\x8b":          # server sent it gzip-compressed
        data = gzip.decompress(data)
    return data


def sql(statement, params=None, show=False):
    with ENGINE.begin() as con:
        res = con.execute(text(statement), params or {})
        if show and res.returns_rows:
            df = pd.DataFrame(res.fetchall(), columns=list(res.keys()))
            print(df.to_string(index=False) if len(df) else "  (no rows)")
            return df
        return res.rowcount


def load_staging(table, df):
    sql(f"TRUNCATE {table}")
    df = df.astype(str).replace({"nan": "", "None": "", "NaT": ""})
    df.to_sql(table, ENGINE, if_exists="append", index=False, method="multi", chunksize=2000)
    print(f"  {table}: {len(df)} rows")


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def as_date(v):
    """Excel date cell -> 'YYYY-MM-DD' (handles real dates and typed text like 17/10/2026)."""
    if v is None or v == "":
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return ""


def as_goals(v):
    if v is None or str(v).strip() == "":
        return ""
    return int(float(v))


def check_label(label):
    if not label.startswith(SEASON) or ("after GW" not in label and "preseason" not in label):
        sys.exit(f'Label "{label}" should look like "{SEASON} after GW6".')


def check_model_file():
    if not os.path.exists(MODEL_PATH):
        sys.exit(f"Model not found: {MODEL_PATH}\nFix MODEL_PATH at the top of weekly.py.")


# ---------------------------------------------------------------- part 1
def make_xg(raw):
    data = json.loads(raw)
    if "dates" not in data:
        raise ValueError("Understat file has no 'dates' section")
    rows = []
    for m in data["dates"]:
        if not m.get("isResult"):
            continue
        h = UNDERSTAT_NAMES.get(m["h"]["title"], m["h"]["title"])
        a = UNDERSTAT_NAMES.get(m["a"]["title"], m["a"]["title"])
        rows.append([SEASON, h, a, round(float(m["xG"]["h"]), 3), round(float(m["xG"]["a"]), 3)])
    write_csv("xg.csv", ["season", "home_team", "away_team", "xg_home", "xg_away"], rows)
    print(f"xg.csv: {len(rows)} played matches")
    return len(rows)


def make_odds(raw):
    text_ = raw.decode("utf-8-sig", errors="replace")
    recs = [r for r in csv.DictReader(text_.splitlines()) if r.get("HomeTeam")]
    if not recs:
        raise ValueError("odds file has no matches (site may have sent a web page)")
    cols = ("AvgH", "AvgD", "AvgA") if "AvgH" in recs[0] else ("B365H", "B365D", "B365A")
    rows = [[SEASON, r["HomeTeam"].strip(), r["AwayTeam"].strip(), r[cols[0]], r[cols[1]], r[cols[2]]]
            for r in recs if r.get(cols[0])]
    write_csv("odds.csv", ["season", "home_team", "away_team", "odds_h", "odds_d", "odds_a"], rows)
    print(f"odds.csv: {len(rows)} matches with odds ({cols[0][:-1]} prices)")
    return len(rows)


def part1(label):
    check_model_file()
    folder = os.path.join(EXPORT_ROOT, label)
    os.makedirs(folder, exist_ok=True)
    os.chdir(folder)
    print(f"Working in {folder}")

    # 1. results from the model's Results sheet
    wb = load_workbook(MODEL_PATH, data_only=True, read_only=True)
    rows = []
    for r in wb["Results"].iter_rows(min_row=2, max_row=381, max_col=7, values_only=True):
        rnd, dt, venue, home, away, hg, ag = (list(r) + [None] * 7)[:7]
        if not home or not away:
            continue
        rows.append({"season": SEASON, "round_no": "" if rnd is None else int(float(rnd)),
                     "match_date": as_date(dt), "venue": venue or "",
                     "home_team": str(home).strip(), "away_team": str(away).strip(),
                     "home_goals": as_goals(hg), "away_goals": as_goals(ag)})
    wb.close()
    results = pd.DataFrame(rows)
    results.to_csv("results.csv", index=False)
    played = int((results["home_goals"] != "").sum())
    print(f"results.csv: {len(results)} fixtures (expect 380), {played} with scores")
    if len(results) != 380:
        sys.exit("Results sheet should have 380 fixtures in rows 2-381. Check the sheet, save, rerun.")

    # 2. xG from Understat (or a hand-saved understat_2026-27.json in this folder)
    json_name = f"understat_{SEASON}.json"
    try:
        raw = fetch(UNDERSTAT_URL, {"X-Requested-With": "XMLHttpRequest",
                                    "Referer": "https://understat.com/league/EPL/2026"})
        json.loads(raw)
        open(json_name, "wb").write(raw)
        print("Understat downloaded")
    except Exception as e:
        print(f"Understat download failed ({e}); using {json_name} in this folder if present")
    if not os.path.exists(json_name):
        sys.exit(f"Save the Understat file (weekly guide Stage B) into this folder as {json_name}, then rerun part1.")
    n_xg = make_xg(open(json_name, "rb").read())

    # 3. odds from football-data (optional)
    have_odds = False
    try:
        raw = fetch(ODDS_URL)
        open("E0.csv", "wb").write(raw)
        have_odds = make_odds(raw) > 0
    except Exception as e:
        print(f"Odds skipped this week ({e})")

    # 4. staging
    print("Loading staging tables")
    load_staging("stg_results", results)
    load_staging("stg_xg", pd.read_csv("xg.csv", dtype=str))
    if have_odds:
        load_staging("stg_odds", pd.read_csv("odds.csv", dtype=str))

    odds_names = ("UNION SELECT TRIM(home_team) FROM stg_odds UNION SELECT TRIM(away_team) FROM stg_odds"
                  if have_odds else "")
    print("\nUnknown team names (should be no rows):")
    unknown = sql(f"""SELECT DISTINCT n AS unknown_name FROM (
                        SELECT TRIM(home_team) n FROM stg_results UNION SELECT TRIM(away_team) FROM stg_results
                        UNION SELECT TRIM(home_team) FROM stg_xg UNION SELECT TRIM(away_team) FROM stg_xg
                        {odds_names}
                      ) x WHERE n NOT IN (SELECT alias FROM team_aliases)""", show=True)
    if len(unknown):
        sys.exit("Add each unknown name to team_aliases, e.g.\n"
                 "  INSERT INTO team_aliases (alias, team_id) SELECT 'Hull', team_id FROM teams WHERE team_name = 'Hull City';\n"
                 "then rerun part1 (nothing has been written to matches yet).")

    # 5. scores, xG, odds
    n = sql("""UPDATE matches m
               SET home_goals = NULLIF(r.home_goals, '')::INT, away_goals = NULLIF(r.away_goals, '')::INT,
                   match_date = COALESCE(NULLIF(r.match_date, '')::DATE, m.match_date)
               FROM stg_results r
               JOIN seasons s ON s.label = r.season
               JOIN team_aliases h ON h.alias = TRIM(r.home_team)
               JOIN team_aliases a ON a.alias = TRIM(r.away_team)
               WHERE m.season_id = s.season_id AND m.home_team_id = h.team_id AND m.away_team_id = a.team_id""")
    print(f"\nUPDATE matches: {n} (expect 380)")
    n = sql("""INSERT INTO match_xg (match_id, xg_home, xg_away)
               SELECT m.match_id, NULLIF(x.xg_home, '')::NUMERIC, NULLIF(x.xg_away, '')::NUMERIC
               FROM stg_xg x JOIN seasons s ON s.label = x.season
               JOIN team_aliases h ON h.alias = TRIM(x.home_team)
               JOIN team_aliases a ON a.alias = TRIM(x.away_team)
               JOIN matches m ON m.season_id = s.season_id AND m.home_team_id = h.team_id AND m.away_team_id = a.team_id
               WHERE NULLIF(x.xg_home, '') IS NOT NULL
               ON CONFLICT (match_id) DO UPDATE SET xg_home = EXCLUDED.xg_home, xg_away = EXCLUDED.xg_away""")
    print(f"match_xg rows written: {n} (expect {played})")
    if have_odds:
        n = sql("""INSERT INTO external_predictions (source, run_label, match_id, p_home, p_draw, p_away)
                   SELECT 'bookmaker', 'odds', m.match_id,
                          ROUND((1/o.h)/(1/o.h+1/o.d+1/o.a), 4), ROUND((1/o.d)/(1/o.h+1/o.d+1/o.a), 4),
                          ROUND((1/o.a)/(1/o.h+1/o.d+1/o.a), 4)
                   FROM (SELECT season, TRIM(home_team) home_team, TRIM(away_team) away_team,
                                NULLIF(odds_h,'')::NUMERIC h, NULLIF(odds_d,'')::NUMERIC d, NULLIF(odds_a,'')::NUMERIC a
                         FROM stg_odds) o
                   JOIN seasons s ON s.label = o.season
                   JOIN team_aliases h ON h.alias = o.home_team JOIN team_aliases a ON a.alias = o.away_team
                   JOIN matches m ON m.season_id = s.season_id AND m.home_team_id = h.team_id AND m.away_team_id = a.team_id
                   WHERE o.h IS NOT NULL AND o.d IS NOT NULL AND o.a IS NOT NULL
                   ON CONFLICT (source, run_label, match_id)
                   DO UPDATE SET p_home = EXCLUDED.p_home, p_draw = EXCLUDED.p_draw, p_away = EXCLUDED.p_away""")
        print(f"bookmaker rows written: {n} (expect about {played})")
        try:
            settle_values("E0.csv")
        except Exception as e:
            print(f"Value log not settled ({type(e).__name__}: {e})")

    # 6. availability (Upgrade 14): same as script 28 blocks 14.5b-14.5d
    print("\n--- availability")
    try:
        import fpl_availability
        sys.argv = ["fpl_availability.py", label]
        fpl_availability.main()
        avail = pd.read_csv("availability.csv", dtype=str, encoding="utf-8-sig", keep_default_na=False)
        load_staging("stg_availability", avail)
        sql("""SELECT DISTINCT TRIM(team) AS unknown_team FROM stg_availability
               WHERE TRIM(team) NOT IN (SELECT alias FROM team_aliases UNION SELECT team_name FROM teams)""", show=True)
        n = sql("""INSERT INTO availability_snapshots
                   SELECT s.run_label, a.team_id, s.missing_att::NUMERIC, s.missing_def::NUMERIC,
                          s.att_factor::NUMERIC, s.def_factor::NUMERIC, s.players_out, s.taken_on::DATE
                   FROM stg_availability s
                   JOIN (SELECT alias AS name, team_id FROM team_aliases
                         UNION SELECT team_name, team_id FROM teams) a ON a.name = TRIM(s.team)
                   ON CONFLICT (run_label, team_id) DO UPDATE
                   SET missing_att = EXCLUDED.missing_att, missing_def = EXCLUDED.missing_def,
                       att_factor = EXCLUDED.att_factor, def_factor = EXCLUDED.def_factor,
                       players_out = EXCLUDED.players_out, taken_on = EXCLUDED.taken_on""")
        print(f"availability snapshot rows: {n} (expect 20)")
    except Exception as e:
        print(f"Availability skipped ({e}). Do it by hand: python fpl_availability.py \"{label}\", "
              "then script 28 blocks 14.5b-14.5d.")

    # 7. checks
    print("\nCheck 1: played = with_xg")
    sql("""SELECT s.label AS season, COUNT(*) AS played, COUNT(x.match_id) AS with_xg FROM matches m
           JOIN seasons s USING (season_id) LEFT JOIN match_xg x USING (match_id)
           WHERE m.home_goals IS NOT NULL AND s.label = :s GROUP BY s.label""", {"s": SEASON}, show=True)
    print("\nCheck 2: games played per team (expect one row: 6 | 20, unless a match was postponed)")
    sql("""SELECT played, COUNT(*) AS teams FROM v_league_table WHERE season = :s
           GROUP BY played ORDER BY played""", {"s": SEASON}, show=True)
    print(f"\nNext, in Excel: open the model, paste {folder}\\availability.csv into the Availability sheet, "
          "set NextGW, refresh the xG table, press F9, wait, Ctrl+S, close. Then run part2.")


# ---------------------------------------------------------------- part 2
def part2(label):
    check_model_file()
    with ENGINE.begin() as con:
        if con.execute(text("SELECT 1 FROM model_runs WHERE run_label = :r"), {"r": label}).first():
            sys.exit(f'"{label}" is already loaded. Use Troubleshooting F1 in the weekly guide to replace it.')

    wb = load_workbook(MODEL_PATH, data_only=True, read_only=True)
    ratings = []
    for r in wb["Team Strength updated"].iter_rows(min_row=2, max_row=21, max_col=11, values_only=True):
        ratings.append({"run_label": label, "team": str(r[0]).strip() if r[0] else "", "att": r[9], "def": r[10]})
    ratings = pd.DataFrame(ratings)
    grid = list(wb["Simulations updated"].iter_rows(min_row=1, max_row=1002, max_col=21, values_only=True))
    wb.close()
    if (ratings["team"] == "").any() or ratings[["att", "def"]].isna().any().any():
        sys.exit("Team Strength updated A/J/K has blanks: refresh xG and press F9 in Excel, save, close, rerun.")

    teams = [str(t).strip() for t in grid[0][1:21]]
    body = [row for row in grid[2:] if row[0] is not None]
    try:
        sims = pd.DataFrame([{"run_label": label, "sim_no": i + 1, "team": t, "position": int(row[j + 1])}
                             for i, row in enumerate(body) for j, t in enumerate(teams)])
    except (TypeError, ValueError):
        sys.exit("Simulations updated has blank or text cells: press F9 in Excel, wait, save, close, rerun.")
    bad = int(sims.groupby("sim_no")["position"].sum().ne(210).sum())
    print(f"ratings: {len(ratings)} teams; sims: {len(body)} seasons; seasons not summing to 210: {bad}")
    if len(body) != 1000 or bad:
        sys.exit("Simulations look wrong: press F9 in Excel (data tables), save, close, rerun.")

    folder = os.path.join(EXPORT_ROOT, label)
    os.makedirs(folder, exist_ok=True)
    ratings.to_csv(os.path.join(folder, "ratings.csv"), index=False)
    sims.to_csv(os.path.join(folder, "sims.csv"), index=False)

    load_staging("stg_ratings", ratings)
    load_staging("stg_sims", sims)
    unknown = sql("""SELECT DISTINCT TRIM(team) AS unknown_name FROM stg_ratings
                     WHERE TRIM(team) NOT IN (SELECT alias FROM team_aliases)
                     UNION SELECT DISTINCT TRIM(team) FROM stg_sims
                     WHERE TRIM(team) NOT IN (SELECT alias FROM team_aliases)""", show=True)
    if len(unknown):
        sys.exit("Add the names above to team_aliases, then rerun part2 (nothing was loaded).")

    with ENGINE.begin() as con:   # all three inserts succeed together or not at all
        con.execute(text("""INSERT INTO model_runs (season_id, run_label)
                            SELECT season_id, :r FROM seasons WHERE label = :s"""), {"r": label, "s": SEASON})
        n1 = con.execute(text("""INSERT INTO team_ratings (run_id, team_id, att, def)
                                 SELECT m.run_id, a.team_id, r.att::NUMERIC, r.def::NUMERIC
                                 FROM stg_ratings r JOIN model_runs m ON m.run_label = r.run_label
                                 JOIN team_aliases a ON a.alias = TRIM(r.team)""")).rowcount
        n2 = con.execute(text("""INSERT INTO sim_results (run_id, sim_no, team_id, final_position)
                                 SELECT m.run_id, s.sim_no::INT, a.team_id, s.position::INT
                                 FROM stg_sims s JOIN model_runs m ON m.run_label = s.run_label
                                 JOIN team_aliases a ON a.alias = TRIM(s.team)""")).rowcount
        if n1 != 20 or n2 != 20000:
            raise SystemExit(f"ratings {n1} / sims {n2} (expected 20 / 20000): nothing saved. Send me this output.")
    print(f"run created; ratings {n1} (expect 20); sims {n2} (expect 20000)")

    # the four Python models, each for this run only
    for mod_name in ("dc_model", "benchmarks", "market_model", "dynamic_sims"):
        print(f"\n--- {mod_name}")
        sys.argv = [f"{mod_name}.py"]          # stop scripts reading weekly.py's own arguments
        try:
            mod = __import__(mod_name)
            mod.ONLY_RUN = label
            mod.main()
            print(f"{mod_name}: done")
        except SystemExit as e:
            print(f"{mod_name}: stopped ({e})")
        except ModuleNotFoundError as e:
            print(f"{mod_name}: not found ({e}), skipped")
        except Exception as e:
            print(f"{mod_name}: FAILED ({type(e).__name__}: {e}). Run it by hand (weekly guide Stage F).")

    print("\n--- ensemble")
    try:
        n = sql("""INSERT INTO external_predictions (source, run_label, match_id, p_home, p_draw, p_away)
                   SELECT 'ens_mkt_50', m.run_label, m.match_id,
                          ROUND(0.5*m.p_home+0.5*k.p_home, 4), ROUND(0.5*m.p_draw+0.5*k.p_draw, 4),
                          ROUND(0.5*m.p_away+0.5*k.p_away, 4)
                   FROM v_main_predictions m
                   JOIN external_predictions k ON k.source = 'market' AND k.run_label = m.run_label AND k.match_id = m.match_id
                   WHERE m.run_label = :r
                   ON CONFLICT (source, run_label, match_id)
                   DO UPDATE SET p_home = EXCLUDED.p_home, p_draw = EXCLUDED.p_draw, p_away = EXCLUDED.p_away""",
                {"r": label})
        print(f"ens_mkt_50 rows: {n}")
    except Exception as e:
        print(f"ensemble skipped ({type(e).__name__}: {e})")

    print("\nRows saved for this run:")
    sql("""SELECT source, COUNT(*) AS rows FROM external_predictions WHERE run_label = :r
           GROUP BY source ORDER BY source""", {"r": label}, show=True)
    print("\nNew forecast (top 6):")
    sql("""SELECT team_name, title_pct, top4_pct, relegation_pct FROM v_sim_summary
           WHERE run_label = :r ORDER BY title_pct DESC LIMIT 6""", {"r": label}, show=True)
    print(f'\nDone. Next: python weekly.py part3 "{label}"')


# ---------------------------------------------------------------- upcoming odds (Upgrade 17)
NEXT_ODDS_COLS = {"AvgH": "avg_h", "AvgD": "avg_d", "AvgA": "avg_a", "Avg>2.5": "avg_over25", "Avg<2.5": "avg_under25",
                  "B365H": "b365_h", "B365D": "b365_d", "B365A": "b365_a", "B365>2.5": "b365_over25",
                  "B365<2.5": "b365_under25", "MaxH": "max_h", "MaxD": "max_d", "MaxA": "max_a",
                  "Max>2.5": "max_over25", "Max<2.5": "max_under25"}


def next_odds():
    """Bookmaker odds for upcoming Premier League games into next_odds and v_next_odds."""
    print("\n=== Upcoming odds (football-data.co.uk fixtures.csv)")
    raw = fetch(NEXT_ODDS_URL).decode("utf-8-sig", errors="replace")
    recs = [r for r in csv.DictReader(raw.splitlines()) if r.get("Div") == "E0"]
    if not recs:
        print("No Premier League games in fixtures.csv yet (it is updated on Fridays and Tuesdays).")
        return 0
    with ENGINE.begin() as con:
        alias = dict(con.execute(text("SELECT alias, team_id FROM team_aliases")).fetchall())
        games = {(h, a): mid for mid, h, a in con.execute(text(
            """SELECT m.match_id, m.home_team_id, m.away_team_id FROM matches m JOIN seasons s USING (season_id)
               WHERE s.label = :s"""), {"s": SEASON}).fetchall()}
    rows, unknown = [], set()
    for r in recs:
        h, a = alias.get(r["HomeTeam"].strip()), alias.get(r["AwayTeam"].strip())
        if h is None or a is None:
            unknown.update(n for n, t in ((r["HomeTeam"].strip(), h), (r["AwayTeam"].strip(), a)) if t is None)
            continue
        mid = games.get((h, a))
        if mid is None:
            continue
        row = {"match_id": mid, "home_team_id": h, "away_team_id": a,
               "kickoff": f"{r.get('Date', '')} {r.get('Time', '')}".strip()}
        for src, col in NEXT_ODDS_COLS.items():
            v = (r.get(src) or "").strip()
            row[col] = float(v) if v else None
        rows.append(row)
    if unknown:
        print("Team names not in team_aliases (add them like part1 says): " + ", ".join(sorted(unknown)))
    cols = ["match_id", "home_team_id", "away_team_id", "kickoff"] + list(NEXT_ODDS_COLS.values())
    with ENGINE.begin() as con:
        con.execute(text("CREATE TABLE IF NOT EXISTS next_odds (match_id INT PRIMARY KEY, home_team_id INT, "
                         "away_team_id INT, kickoff TEXT, "
                         + ", ".join(f"{c} NUMERIC" for c in NEXT_ODDS_COLS.values())
                         + ", downloaded_at TIMESTAMP DEFAULT now())"))
        for row in rows:
            con.execute(text(f"""INSERT INTO next_odds ({', '.join(cols)}, downloaded_at)
                                 VALUES ({', '.join(':' + c for c in cols)}, now())
                                 ON CONFLICT (match_id) DO UPDATE SET """
                             + ", ".join(f"{c} = EXCLUDED.{c}" for c in cols[1:]) + ", downloaded_at = now()"), row)
        con.execute(text("""CREATE OR REPLACE VIEW v_next_odds AS
            SELECT ht.team_name || ' v ' || at.team_name AS fixture, ht.team_name AS home_team,
                   at.team_name AS away_team, m.round_no AS gameweek, o.kickoff,
                   o.avg_h, o.avg_d, o.avg_a, o.avg_over25, o.avg_under25,
                   o.b365_h, o.b365_d, o.b365_a, o.b365_over25, o.b365_under25,
                   o.max_h, o.max_d, o.max_a, o.max_over25, o.max_under25, o.downloaded_at
            FROM next_odds o
            JOIN matches m ON m.match_id = o.match_id
            JOIN teams ht ON ht.team_id = o.home_team_id
            JOIN teams at ON at.team_id = o.away_team_id
            WHERE m.home_goals IS NULL"""))
    print(f"next_odds: {len(rows)} upcoming games with odds (v_next_odds shows the unplayed ones)")
    return len(rows)


# ---------------------------------------------------------------- value log (Upgrade 17)
VALUE_MARKETS = ["home", "draw", "away", "over25", "under25"]
OPEN_COLS = {"b365": ["b365_h", "b365_d", "b365_a", "b365_over25", "b365_under25"],
             "avg": ["avg_h", "avg_d", "avg_a", "avg_over25", "avg_under25"],
             "max": ["max_h", "max_d", "max_a", "max_over25", "max_under25"]}
CLOSE_COLS = {"b365": ["B365CH", "B365CD", "B365CA", "B365C>2.5", "B365C<2.5"],
              "avg": ["AvgCH", "AvgCD", "AvgCA", "AvgC>2.5", "AvgC<2.5"],
              "max": ["MaxCH", "MaxCD", "MaxCA", "MaxC>2.5", "MaxC<2.5"]}
VALUE_DDL = """CREATE TABLE IF NOT EXISTS value_log (
    match_id INT, market TEXT, run_label TEXT, source TEXT,
    model_prob NUMERIC, fair_prob NUMERIC, blended_prob NUMERIC, taken_odds NUMERIC, edge NUMERIC,
    flag TEXT, logged_at TIMESTAMP DEFAULT now(),
    closing_odds NUMERIC, closing_fair_prob NUMERIC, won BOOLEAN, settled_at TIMESTAMP,
    PRIMARY KEY (match_id, market))"""


def poisson_markets(lh, la):
    """Main-model chances for home / draw / away / over 2.5 / under 2.5 from the two expected goals."""
    from math import exp, factorial
    ph = [exp(-lh) * lh ** k / factorial(k) for k in range(11)]
    pa = [exp(-la) * la ** k / factorial(k) for k in range(11)]
    h = d = a = over = 0.0
    for i in range(11):
        for j in range(11):
            q = ph[i] * pa[j]
            if i > j:
                h += q
            elif i == j:
                d += q
            else:
                a += q
            if i + j >= 3:
                over += q
    tot = h + d + a
    return [h / tot, d / tot, a / tot, over / tot, 1 - over / tot]


def fair(odds):
    """Bookie's chances with the margin taken out (None if any price is missing)."""
    if any(o is None or o <= 1 for o in odds):
        return [None] * len(odds)
    imp = [1 / o for o in odds]
    return [i / sum(imp) for i in imp]


def log_values():
    """Log every next-gameweek price with the model's chance, the blended chance and the edge."""
    cols = OPEN_COLS[VALUE_SOURCE]
    with ENGINE.begin() as con:
        con.execute(text(VALUE_DDL))
        run = con.execute(text("""SELECT run_label FROM model_runs m JOIN seasons s USING (season_id)
                                  WHERE s.label = :s AND run_label LIKE '%after GW%'
                                  ORDER BY run_id DESC LIMIT 1"""), {"s": SEASON}).scalar()
        games = con.execute(text(f"""SELECT o.match_id, l.lambda_home, l.lambda_away, {', '.join('o.' + c for c in cols)}
                                     FROM next_odds o
                                     JOIN matches m ON m.match_id = o.match_id AND m.home_goals IS NULL
                                     JOIN match_lambdas l ON l.match_id = o.match_id AND l.run_label = :r"""),
                            {"r": run}).fetchall()
        n = flagged = 0
        for g in games:
            mid, lh, la, odds = g[0], float(g[1]), float(g[2]), [None if v is None else float(v) for v in g[3:]]
            model = poisson_markets(lh, la)
            fp = fair(odds[:3]) + fair(odds[3:])
            for k, mk in enumerate(VALUE_MARKETS):
                if odds[k] is None or fp[k] is None:
                    continue
                blended = VALUE_MODEL_WEIGHT * model[k] + (1 - VALUE_MODEL_WEIGHT) * fp[k]
                edge = blended * odds[k] - 1
                flag = "VALUE" if edge >= VALUE_FLAG_EDGE else ("small" if edge >= 0 else "no")
                con.execute(text("""INSERT INTO value_log (match_id, market, run_label, source, model_prob, fair_prob,
                                        blended_prob, taken_odds, edge, flag, logged_at)
                                    VALUES (:m, :k, :r, :src, :mp, :fp, :bp, :o, :e, :f, now())
                                    ON CONFLICT (match_id, market) DO UPDATE SET run_label = EXCLUDED.run_label,
                                        source = EXCLUDED.source, model_prob = EXCLUDED.model_prob,
                                        fair_prob = EXCLUDED.fair_prob, blended_prob = EXCLUDED.blended_prob,
                                        taken_odds = EXCLUDED.taken_odds, edge = EXCLUDED.edge,
                                        flag = EXCLUDED.flag, logged_at = now()
                                    WHERE value_log.settled_at IS NULL"""),
                            {"m": mid, "k": mk, "r": run, "src": VALUE_SOURCE, "mp": round(model[k], 4),
                             "fp": round(fp[k], 4), "bp": round(blended, 4), "o": odds[k], "e": round(edge, 4),
                             "f": flag})
                n += 1
                flagged += flag == "VALUE"
        make_value_views(con)
    print(f"value_log: {n} prices logged from {len(games)} games ({VALUE_SOURCE} odds, run {run}); "
          f"{flagged} flagged VALUE at +{VALUE_FLAG_EDGE:.0%} or more")


def settle_values(e0_path):
    """After the games: closing odds (football-data C columns) and the result for every logged price."""
    if not os.path.exists(e0_path):
        return
    with ENGINE.begin() as con:
        if not con.execute(text("SELECT to_regclass('value_log')")).scalar():
            return
        alias = dict(con.execute(text("SELECT alias, team_id FROM team_aliases")).fetchall())
        open_ = {(h, a): (mid, hg, ag) for mid, h, a, hg, ag in con.execute(text(
            """SELECT m.match_id, m.home_team_id, m.away_team_id, m.home_goals, m.away_goals
               FROM matches m JOIN value_log v USING (match_id)
               WHERE v.settled_at IS NULL AND m.home_goals IS NOT NULL""")).fetchall()}
        if not open_:
            print("value_log: nothing waiting to settle")
            return
        recs = list(csv.DictReader(open(e0_path, encoding="utf-8-sig", errors="replace")))
        n = 0
        for r in recs:
            key = (alias.get((r.get("HomeTeam") or "").strip()), alias.get((r.get("AwayTeam") or "").strip()))
            if key not in open_:
                continue
            mid, hg, ag = open_[key]
            src = con.execute(text("SELECT MAX(source) FROM value_log WHERE match_id = :m"), {"m": mid}).scalar()
            close = []
            for c in CLOSE_COLS.get(src, CLOSE_COLS["b365"]):
                v = (r.get(c) or "").strip()
                close.append(float(v) if v else None)
            cf = fair(close[:3]) + fair(close[3:])
            won = {"home": hg > ag, "draw": hg == ag, "away": hg < ag,
                   "over25": hg + ag >= 3, "under25": hg + ag <= 2}
            for k, mk in enumerate(VALUE_MARKETS):
                n += con.execute(text("""UPDATE value_log SET closing_odds = :c, closing_fair_prob = :cf,
                                             won = :w, settled_at = now()
                                         WHERE match_id = :m AND market = :k AND settled_at IS NULL"""),
                                 {"c": close[k], "cf": None if cf[k] is None else round(cf[k], 4),
                                  "w": bool(won[mk]), "m": mid, "k": mk}).rowcount
        make_value_views(con)
    print(f"value_log: {n} prices settled with closing odds and results")


def make_value_views(con):
    con.execute(text("""CREATE OR REPLACE VIEW v_value_log AS
        SELECT v.run_label, m.round_no AS gameweek, ht.team_name || ' v ' || at.team_name AS fixture,
               CASE v.market WHEN 'home' THEN ht.team_name || ' win' WHEN 'away' THEN at.team_name || ' win'
                    WHEN 'draw' THEN 'Draw' WHEN 'over25' THEN 'Over 2.5 goals' ELSE 'Under 2.5 goals' END AS bet,
               v.source, v.flag,
               ROUND(100 * v.model_prob, 1) AS model_pct, ROUND(100 * v.fair_prob, 1) AS bookie_fair_pct,
               ROUND(100 * v.blended_prob, 1) AS blended_pct, v.taken_odds, ROUND(100 * v.edge, 1) AS edge_pct,
               v.closing_odds,
               ROUND(100 * (v.taken_odds / NULLIF(v.closing_odds, 0) - 1), 1) AS clv_pct,
               v.won,
               CASE WHEN v.won IS NULL THEN NULL WHEN v.won THEN v.taken_odds - 1 ELSE -1 END AS profit_1u,
               v.logged_at, v.settled_at
        FROM value_log v
        JOIN matches m ON m.match_id = v.match_id
        JOIN teams ht ON ht.team_id = m.home_team_id
        JOIN teams at ON at.team_id = m.away_team_id"""))
    con.execute(text("""CREATE OR REPLACE VIEW v_value_summary AS
        SELECT g.grp AS flag_group,
               COUNT(*) AS prices_logged,
               COUNT(*) FILTER (WHERE v.settled_at IS NOT NULL) AS settled,
               ROUND(AVG(100 * v.edge), 1) AS avg_edge_pct,
               ROUND(AVG(100 * (v.taken_odds / NULLIF(v.closing_odds, 0) - 1))
                     FILTER (WHERE v.closing_odds IS NOT NULL), 1) AS avg_clv_pct,
               ROUND(100.0 * AVG((v.taken_odds > v.closing_odds)::INT)
                     FILTER (WHERE v.closing_odds IS NOT NULL), 1) AS beat_close_pct,
               ROUND(100.0 * AVG(v.won::INT) FILTER (WHERE v.won IS NOT NULL), 1) AS won_pct,
               ROUND(100.0 * AVG(1 / v.taken_odds) FILTER (WHERE v.won IS NOT NULL), 1) AS odds_implied_pct,
               ROUND(SUM(CASE WHEN v.won THEN v.taken_odds - 1 WHEN NOT v.won THEN -1 END), 2) AS profit_1u,
               ROUND(100 * SUM(CASE WHEN v.won THEN v.taken_odds - 1 WHEN NOT v.won THEN -1 END)
                     / NULLIF(COUNT(*) FILTER (WHERE v.won IS NOT NULL), 0), 1) AS roi_pct
        FROM value_log v
        CROSS JOIN LATERAL (VALUES (v.flag), ('All prices')) g(grp)
        GROUP BY g.grp
        ORDER BY CASE g.grp WHEN 'VALUE' THEN 1 WHEN 'small' THEN 2 WHEN 'no' THEN 3 ELSE 4 END"""))


# ---------------------------------------------------------------- part 3 (Upgrade 17)
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
NEXT_SCRIPTS = ["player_minutes.py", "player_rates.py", "team_counts.py", "player_markets.py", "match_markets.py"]


def gw_of(label):
    try:
        return int(label.split("GW")[1])
    except (IndexError, ValueError):
        sys.exit(f'part3 needs an "after GW" label, e.g. "{SEASON} after GW7".')


def run_script(*args):
    """Run one of the project scripts in its own process (own arguments, own memory) and stop on failure."""
    print(f"\n--- python {' '.join(args)}", flush=True)
    r = subprocess.run([sys.executable, *args], cwd=SCRIPT_DIR)
    if r.returncode != 0:
        sys.exit(f"\n{args[0]} stopped with an error (above). Fix it, then rerun part3: "
                 "finished steps simply replace their own rows.")


def referees_ready(next_gw, skip):
    """True when referees_next.csv is for the next gameweek and filled in; otherwise writes it and stops."""
    path = os.path.join(SCRIPT_DIR, "referees_next.csv")
    if skip:
        print("Referees skipped (--skip-referees): factor 1 and 'referee TBC' for this gameweek.")
        return False
    df = pd.read_csv(path, encoding="utf-8-sig") if os.path.exists(path) else None
    if df is None or len(df) == 0 or int(df["gameweek"].iloc[0]) != next_gw:
        run_script("player_data.py", "referees")
        df = pd.read_csv(path, encoding="utf-8-sig") if os.path.exists(path) else None
        if df is None or int(df["gameweek"].iloc[0]) != next_gw:
            sys.exit(f"Could not make referees_next.csv for GW{next_gw} (see above). "
                     "Rerun with --skip-referees to carry on without them.")
    blank = df["referee"].isna() | (df["referee"].astype(str).str.strip() == "")
    if blank.any():
        print(f"\nreferees_next.csv (in {SCRIPT_DIR}) needs the referee for these GW{next_gw} games:")
        for r in df[blank].itertuples(index=False):
            print(f"   {r.home_team} v {r.away_team}")
        sys.exit(f"\nOpen it with:  notepad \"{path}\"\n"
                 f"type the names from premierleague.com 'Match officials for Matchweek {next_gw}' "
                 "(e.g. Michael Oliver or M Oliver), save, then run part3 again.")
    return True


def part3(label, skip_referees=False):
    gw = gw_of(label)
    with ENGINE.begin() as con:
        if not con.execute(text("SELECT 1 FROM model_runs WHERE run_label = :r"), {"r": label}).first():
            sys.exit(f'"{label}" is not loaded yet: run part1 and part2 first.')

    # 1. referees for the next gameweek (stops here the first time so you can fill them in)
    print(f"=== 1. GW{gw + 1} referees")
    if referees_ready(gw + 1, skip_referees):
        run_script("player_data.py", "referees", "load")

    # 2. 15.5a: each team's expected goals for the next gameweek, from the main model (this run only)
    print("\n=== 2. Expected goals per fixture (15.5a)")
    n = sql("""INSERT INTO match_lambdas (run_label, match_id, lambda_home, lambda_away)
               SELECT r.run_label, m.match_id,
                      h.att * a.def * ra.avg_goals * v.home_factor,
                      a.att * h.def * ra.avg_goals * v.away_factor
               FROM v_runs r
               JOIN v_run_avg ra ON ra.run_id = r.run_id
               JOIN model_variants v ON v.variant = CASE WHEN r.season >= '2026-27' AND r.gameweek >= 6
                                                         THEN 'home_adv' ELSE 'value_prom_50' END
               JOIN matches m ON m.season_id = r.season_id AND m.round_no = r.gameweek + 1
               JOIN v_variant_ratings h ON h.variant = v.variant AND h.run_id = r.run_id AND h.team_id = m.home_team_id
               JOIN v_variant_ratings a ON a.variant = v.variant AND a.run_id = r.run_id AND a.team_id = m.away_team_id
               WHERE r.run_label = :r
               ON CONFLICT (run_label, match_id) DO UPDATE
               SET lambda_home = EXCLUDED.lambda_home, lambda_away = EXCLUDED.lambda_away""", {"r": label})
    print(f"match_lambdas rows written: {n}")

    # 3. player and team data for the season so far
    print("\n=== 3. Player and team data")
    run_script("player_data.py", "update")

    # 4. the five 'next' scripts, in order (each uses the one before)
    print("\n=== 4. Next gameweek markets")
    for s_ in NEXT_SCRIPTS:
        run_script(s_, "next")

    # 5. bookmaker odds for the upcoming games (for the H2H value section)
    try:
        if next_odds():
            log_values()
    except Exception as e:
        print(f"Upcoming odds skipped ({type(e).__name__}: {e}); run python weekly.py odds later.")

    # 6. checks
    check(label)
    print(f"\nDone. In the dashboard: add '{label}' at the bottom of Lists column A, pick it in H2H C6, "
          "then press Refresh All (with the model closed).")


# ---------------------------------------------------------------- checks
def check(label):
    gw = gw_of(label)
    print(f"\n=== Checks for {label} (predicting GW{gw + 1})")
    problems = []

    def one(sql_text, params=None):
        with ENGINE.begin() as con:
            return con.execute(text(sql_text), params or {}).scalar()

    def mark(ok, msg):
        print(f"  {'OK   ' if ok else 'CHECK'}  {msg}")
        if not ok:
            problems.append(msg)

    p = {"r": label, "s": SEASON, "g": gw, "n": gw + 1}
    try:
        fixtures = one("""SELECT COUNT(*) FROM matches m JOIN seasons s USING (season_id)
                          WHERE s.label = :s AND m.round_no = :n""", p) or 0
        played = one("""SELECT COUNT(*) FROM matches m JOIN seasons s USING (season_id)
                        WHERE s.label = :s AND m.home_goals IS NOT NULL""", p)
        with_xg = one("""SELECT COUNT(*) FROM matches m JOIN seasons s USING (season_id) JOIN match_xg x USING (match_id)
                         WHERE s.label = :s AND m.home_goals IS NOT NULL""", p)
        mark(played == with_xg, f"part1: {played} matches played, {with_xg} with xG")
        ratings = one("""SELECT COUNT(*) FROM team_ratings t JOIN model_runs m USING (run_id) WHERE m.run_label = :r""", p)
        sims = one("""SELECT COUNT(*) FROM sim_results t JOIN model_runs m USING (run_id) WHERE m.run_label = :r""", p)
        mark(ratings == 20 and sims == 20000, f"part2: run has {ratings} ratings (20) and {sims} sims (20000)")
        for src in ("dixon_coles", "market", "ens_mkt_50", "pi_ratings", "dc_xg"):
            k = one("SELECT COUNT(*) FROM external_predictions WHERE source = :x AND run_label = :r", {**p, "x": src})
            mark(k == 380, f"part2: {src} {k} predictions (380)")
        av = one("SELECT COUNT(*) FROM availability_snapshots WHERE run_label = :r", p)
        mark(av == 20, f"part1: availability snapshot {av} teams (20)")

        print(f"  GW{gw + 1} has {fixtures} fixtures in the matches table")
        lam = one("SELECT COUNT(*) FROM match_lambdas WHERE run_label = :r", p)
        mark(lam == fixtures and fixtures > 0, f"15.5a: {lam} fixtures with expected goals ({fixtures})")
        refs = one("SELECT COUNT(*) FROM fixture_referees WHERE season = :s AND gameweek = :n "
                   "AND COALESCE(referee, '') <> ''", p)
        known = one("""SELECT COUNT(*) FROM fixture_referees f
                       JOIN v_referee_stats r ON r.season = 'all' AND r.referee = f.referee
                       WHERE f.season = :s AND f.gameweek = :n""", p)
        mark(refs == fixtures, f"referees: {refs} of {fixtures} GW{gw + 1} fixtures have a referee")
        mark(known == refs, f"referees: {known} of {refs} have history (a blank = new referee or a spelling "
                            "that differs from football-data)")
        lo, hi = (one(f"""SELECT {fn}(s) FROM (SELECT SUM(p_start) s FROM player_minutes
                          WHERE mode = 'next' AND run_label = :r GROUP BY match_id, team_id) x""", p)
                  for fn in ("MIN", "MAX"))
        mark(lo is not None and abs(float(lo) - 11) < 0.05 and abs(float(hi) - 11) < 0.05,
             f"player_minutes: expected starters per team {lo and round(float(lo), 2)}-{hi and round(float(hi), 2)} (11)")
        bad = one("""SELECT COUNT(*) FROM player_minutes WHERE mode = 'next' AND run_label = :r
                     AND p_start > availability * 0.97 + 0.001""", p)
        mark(bad == 0, f"player_minutes: {bad} players starting more often than their FPL chance allows (0)")
        tc = one("SELECT COUNT(*) FROM team_counts WHERE mode = 'next' AND run_label = :r", p)
        mark(tc == fixtures * 21, f"team_counts: {tc} rows ({fixtures * 21})")
        pmk = one("SELECT COUNT(DISTINCT match_id) FROM player_markets WHERE mode = 'next' AND run_label = :r", p)
        mark(pmk == fixtures, f"player_markets: {pmk} fixtures ({fixtures})")
        mm = one("SELECT COUNT(*) FROM match_markets WHERE mode = 'next' AND run_label = :r", p)
        mark(mm == fixtures * 100, f"match_markets: {mm} rows ({fixtures * 100})")
        fc = one("SELECT COUNT(DISTINCT team) FROM v_fixture_card", p)
        mark(fc == 2 * fixtures, f"dashboard v_fixture_card: {fc} teams ({2 * fixtures})")
        try:
            no = one("SELECT COUNT(*) FROM v_next_odds WHERE gameweek = :n", p)
            print(f"  info   upcoming odds: {no} of {fixtures} GW{gw + 1} games have bookmaker odds "
                  "(fixtures.csv updates Fri/Tue; rerun python weekly.py odds then)")
        except Exception:
            print("  info   upcoming odds: not loaded yet (python weekly.py odds)")
        pc = one("SELECT COUNT(DISTINCT player_key) FROM v_player_card", p)
        mark(pc and pc > 300, f"dashboard v_player_card: {pc} players")
    except Exception as e:
        mark(False, f"check stopped early: {type(e).__name__}: {e}")
    print(f"\n{'All checks OK.' if not problems else str(len(problems)) + ' to look at (CHECK lines above).'}")
    return problems


if __name__ == "__main__":
    cmds = ("part1", "part2", "part3", "check")
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if args == ["odds"]:
        if next_odds():
            log_values()
        sys.exit(0)
    if len(args) != 2 or args[0] not in cmds:
        sys.exit('Usage: python weekly.py part1|part2|part3|check "2026-27 after GW7"   or   python weekly.py odds')
    check_label(args[1])
    if args[0] == "part3":
        part3(args[1], skip_referees="--skip-referees" in sys.argv)
    elif args[0] == "check":
        check(args[1])
    else:
        {"part1": part1, "part2": part2}[args[0]](args[1])
